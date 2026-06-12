"""Build the sfast-2026-1 golden case (deal.json + expected CSVs) from the
uploaded workbook 'SFAST2026-1 deal config.xlsx'.

Usage: python scripts/extract_sfast_case.py
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pandas as pd

CASES = Path(__file__).resolve().parents[1] / "tests" / "golden" / "cases"
WB = CASES / "SFAST2026-1 deal config.xlsx"
OUT = CASES / "sfast-2026-1"

# sheet column -> engine column (1-indexed sheet columns from the header rows)
COLLAT_COLS = {
    4: "interest",
    7: "end_trust",
    8: "losses",
    9: "chargeoffs",
    13: "sched_prin",
    14: "prepay_prin",
    15: "recoveries",
    20: "adjusted_pool",
}
# per-class blocks: first sheet column of (Principal, Interest, ..., Balance at +3)
BOND_BLOCKS = {"A1": 35, "A2A": 45, "A2B": 55, "A3": 66, "A4": 76, "B": 86, "C": 96}

SCEN_SHEETS = {"scen1": "CF scenario 1", "scen2": "CF scenario 2"}
SOFR_COL = 145  # SOFR_NYF_30D projection, percent


def extract_expected(ws, scen: str) -> None:
    rows = list(ws.iter_rows(values_only=True))
    collat_rows, bond_rows, sofr = [], [], []
    for row in rows[13:]:  # first data row = period 1
        if not isinstance(row[0], (int, float)) or row[0] is None or row[0] < 1:
            break
        period = int(row[0])
        rec = {"period": period}
        for col, name in COLLAT_COLS.items():
            v = row[col - 1]
            rec[name] = float(v) if v not in (None, "", " ") else 0.0
        collat_rows.append(rec)
        sofr.append(float(row[SOFR_COL - 1]) / 100.0)
        for cid, c0 in BOND_BLOCKS.items():
            prin, inte, bal = row[c0 - 1], row[c0], row[c0 + 2]
            bond_rows.append({
                "period": period,
                "class_id": cid,
                "prin_paid": float(prin) if prin is not None else 0.0,
                "interest_paid": float(inte) if inte is not None else 0.0,
                "end_balance": float(bal) if bal is not None else 0.0,
            })
    pd.DataFrame(collat_rows).to_csv(OUT / f"expected_{scen}__collateral.csv", index=False)
    pd.DataFrame(bond_rows).to_csv(OUT / f"expected_{scen}__bonds.csv", index=False)
    (OUT / f"sofr_{scen}.json").write_text(json.dumps(sofr))


def build_deal(wb) -> dict:
    reps = [
        {
            "id": f"R{int(r[0])}",
            "balance": float(r[1]),
            "gross_rate": float(r[2]) / 100.0,
            "original_term": int(r[3]),
            "remaining_term": int(r[4]),
        }
        for r in wb["collateral"].iter_rows(min_row=2, values_only=True)
        if r[0] is not None
    ]

    classes = [
        {"id": "A1", "balance": 250_000_000, "coupon": {"type": "fixed", "rate": 0.03827},
         "day_count": "ACT/360"},
        {"id": "A2A", "balance": 250_000_000, "coupon": {"type": "fixed", "rate": 0.0391},
         "day_count": "30/360"},
        {"id": "A2B", "balance": 250_000_000,
         "coupon": {"type": "floating", "index": "SOFR_NYF_30D", "margin": 0.0046},
         "day_count": "ACT/360"},
        {"id": "A3", "balance": 500_000_000, "coupon": {"type": "fixed", "rate": 0.0396},
         "day_count": "30/360"},
        {"id": "A4", "balance": 119_520_000, "coupon": {"type": "fixed", "rate": 0.0407},
         "day_count": "30/360"},
        {"id": "B", "balance": 71_880_000, "coupon": {"type": "fixed", "rate": 0.0427},
         "day_count": "30/360"},
        {"id": "C", "balance": 58_600_000, "coupon": {"type": "fixed", "rate": 0.0446},
         "day_count": "30/360"},
    ]

    tree = {
        "type": "group", "name": "notes", "mode": "sequential",
        "children": [
            {"type": "group", "name": "classA", "mode": "sequential", "children": [
                {"type": "class", "class_id": "A1"},
                {"type": "group", "name": "A2", "mode": "pro_rata", "children": [
                    {"type": "class", "class_id": "A2A"},
                    {"type": "class", "class_id": "A2B"},
                ]},
                {"type": "class", "class_id": "A3"},
                {"type": "class", "class_id": "A4"},
            ]},
            {"type": "class", "class_id": "B"},
            {"type": "class", "class_id": "C"},
        ],
    }

    adj = "adjusted"
    oc = {"kind": "pct_current_pool", "value": 0.0675,
          "floor_kind": "pct_original_pool", "floor_value": 0.01}

    def steps():
        out = [
            {"id": "svc_fee", "type": "pay_fees", "source": "total_collections",
             "label": "2nd: servicing fee", "fees": ["servicing"]},
            {"id": "svc_fee_rsv", "type": "pay_fees", "source": "reserve:reserve",
             "fees": ["servicing"]},
            {"id": "trustee_fee", "type": "pay_fees", "source": "total_collections",
             "label": "3rd: trustee fees", "fees": ["trustee"]},
            {"id": "trustee_fee_rsv", "type": "pay_fees", "source": "reserve:reserve",
             "fees": ["trustee"]},
            {"id": "int_a", "type": "pay_interest", "source": "total_collections",
             "label": "4th: class A interest", "targets": ["classA"]},
            {"id": "int_a_rsv", "type": "pay_interest", "source": "reserve:reserve",
             "targets": ["classA"]},
            {"id": "first_alloc", "type": "pay_principal", "source": "total_collections",
             "label": "5th: First Allocation of Principal", "amount_rule": "priority_pda",
             "pool_basis": adj, "targets": ["classA"]},
            {"id": "first_alloc_rsv", "type": "pay_principal", "source": "reserve:reserve",
             "amount_rule": "priority_pda", "pool_basis": adj, "targets": ["classA"]},
            {"id": "int_b", "type": "pay_interest", "source": "total_collections",
             "label": "6th: class B interest", "targets": ["B"]},
            {"id": "int_b_rsv", "type": "pay_interest", "source": "reserve:reserve",
             "targets": ["B"]},
            {"id": "second_alloc", "type": "pay_principal", "source": "total_collections",
             "label": "7th: Second Allocation of Principal", "amount_rule": "priority_pda",
             "pool_basis": adj, "targets": ["classA", "B"]},
            {"id": "second_alloc_rsv", "type": "pay_principal", "source": "reserve:reserve",
             "amount_rule": "priority_pda", "pool_basis": adj, "targets": ["classA", "B"]},
            {"id": "int_c", "type": "pay_interest", "source": "total_collections",
             "label": "8th: class C interest", "targets": ["C"]},
            {"id": "int_c_rsv", "type": "pay_interest", "source": "reserve:reserve",
             "targets": ["C"]},
            {"id": "third_alloc", "type": "pay_principal", "source": "total_collections",
             "label": "9th: Third Allocation of Principal", "amount_rule": "priority_pda",
             "pool_basis": adj, "targets": ["classA", "B", "C"]},
            {"id": "third_alloc_rsv", "type": "pay_principal", "source": "reserve:reserve",
             "amount_rule": "priority_pda", "pool_basis": adj,
             "targets": ["classA", "B", "C"]},
            {"id": "fund_rsv", "type": "fund_reserve", "source": "total_collections",
             "label": "10th: replenish reserve", "account": "reserve"},
            {"id": "retire", "type": "retire_bonds", "source": "total_collections",
             "label": "reserve-to-retire option", "reserve": "reserve"},
            {"id": "regular_alloc", "type": "pay_principal", "source": "total_collections",
             "label": "11th: Regular Allocation of Principal", "amount_rule": "regular_pda",
             "pool_basis": adj, "target_oc": oc, "targets": ["notes"]},
            {"id": "residual", "type": "release_residual", "source": "total_collections",
             "label": "13th: to certificateholder", "to": "CERT"},
        ]
        return out

    def scenario(name: str, sofr: list[float], loss: dict) -> dict:
        return {
            "name": name,
            "prepay": {"speed": {"type": "scalar", "value": 0.014},
                       "speed_type": "voluntary", "speed_unit": "abs",
                       "prepay_base": "gross_of_defaults"},
            "loss": loss,
            "recoveries_to": "principal",
            "index_curves": {"SOFR_NYF_30D": {"type": "vector", "values": sofr}},
        }

    sofr1 = json.loads((OUT / "sofr_scen1.json").read_text())
    sofr2 = json.loads((OUT / "sofr_scen2.json").read_text())
    scen1_loss = {
        "defaults": {"type": "cdr", "cdr": {"type": "scalar", "value": 0.0225}},
        "severity": {"type": "scalar", "value": 0.5},
        "charge_off_lag": 3, "recovery_lag": 3, "recovery_lag_from": "default",
        "suppress_defaults_near_maturity": True,
    }
    scen2_loss = {
        # the provided run takes the loss dollars literally off the original
        # balance (realized CNL = input exactly) -> original_MDR convention;
        # the timing curve positions loss recognition (charge-offs)
        "defaults": {"type": "cum_loss", "cum_net_loss": 0.0225,
                     "timing": [40, 35, 20, 5], "timing_unit": "annual",
                     "timing_applies_to": "losses",
                     "method": "original_MDR", "allocation": "pool"},
        "severity": {"type": "scalar", "value": 0.5},
        "charge_off_lag": 3, "recovery_lag": 3, "recovery_lag_from": "default",
        "suppress_defaults_near_maturity": True,
    }

    return {
        "schema_version": 1,
        "id": "sfast-2026-1",
        "name": "SFS Auto Receivables Securitization Trust 2026-1",
        "num_periods": 75,
        "dates": {"closing_date": "2026-02-26", "first_payment_date": "2026-03-20",
                  "business_day_adjust": "following"},
        "collateral": {"asset_class": "amortizing_loan", "replines": reps,
                       "servicing_fee_rate": 0.0, "fee_basis": "trust"},
        "structure": {"classes": classes, "allocation_tree": tree},
        "fees": [
            {"name": "servicing", "rate": 0.01, "basis": "trust"},
            {"name": "trustee", "fixed": 1425.0},
        ],
        "reserve_accounts": [
            {"name": "reserve", "target_kind": "fixed",
             "target_value": 3906253.028725, "initial_balance": 3906253.028725},
        ],
        # stepdown strike applies inside the PDA formulas once A2B retires;
        # the reported adjusted-balance column stays at the primary strike
        "ysoc": {
            "required_rate": 0.063, "stepdown_rate": 0.058,
            "stepdown_when_class_zero": "A2B",
            "initial_amount": 71708995.93, "basis": "trust",
        },
        "waterfall": {"mode": "combined", "waterfalls": [
            {"name": "priority of payments", "steps": steps()},
        ]},
        "scenarios": [scenario("scen1", sofr1, scen1_loss),
                      scenario("scen2", sofr2, scen2_loss)],
    }


def main() -> None:
    OUT.mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(WB, data_only=True)
    for scen, sheet in SCEN_SHEETS.items():
        extract_expected(wb[sheet], scen)
    deal = build_deal(wb)
    (OUT / "deal.json").write_text(json.dumps(deal, indent=2))
    (OUT / "config.json").write_text(json.dumps({"abs_tol": 0.001}, indent=2))
    print(f"wrote case to {OUT}")


if __name__ == "__main__":
    main()
